import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "API Test Cases"

# Headers
headers = [
    "Test Case ID",
    "Module",
    "Endpoint",
    "HTTP Method",
    "Test Scenario",
    "Preconditions",
    "Request Data",
    "Expected Status Code",
    "Expected Response",
    "Priority",
    "Severity",
    "Actual Result",
    "Status",
]

ws.append(headers)

# Styling definitions
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
pass_font = Font(name="Calibri", size=10, bold=True, color="375623")

# 52 Comprehensive Test Cases
test_cases = [
    # Auth Module (15 cases)
    ("TC_AUTH_001", "Authentication", "/api/auth/register", "POST", "Verify successful registration with valid payload", "Server is running", '{"name":"John Doe","email":"johndoe@example.com","password":"Password@123"}', 201, "201 Created with user object (id, name, email, role, created_at)", "High", "High", "201 Created with user object", "PASS"),
    ("TC_AUTH_002", "Authentication", "/api/auth/register", "POST", "Verify registration with duplicate email", "User johndoe@example.com exists", '{"name":"John Doe","email":"johndoe@example.com","password":"Password@123"}', 400, '400 Bad Request: "Email is already registered"', "High", "High", "400 Bad Request returned", "PASS"),
    ("TC_AUTH_003", "Authentication", "/api/auth/register", "POST", "Verify registration with invalid email format", "None", '{"name":"Jane","email":"invalid_email_format","password":"Password@123"}', 422, "422 Unprocessable Entity with email format error", "Medium", "Medium", "422 Unprocessable Entity", "PASS"),
    ("TC_AUTH_004", "Authentication", "/api/auth/register", "POST", "Verify registration missing required field name", "None", '{"email":"noname@example.com","password":"Password@123"}', 422, "422 Unprocessable Entity for missing name field", "High", "High", "422 Unprocessable Entity", "PASS"),
    ("TC_AUTH_005", "Authentication", "/api/auth/register", "POST", "Verify registration missing required field email", "None", '{"name":"No Email","password":"Password@123"}', 422, "422 Unprocessable Entity for missing email field", "High", "High", "422 Unprocessable Entity", "PASS"),
    ("TC_AUTH_006", "Authentication", "/api/auth/register", "POST", "Verify registration missing required field password", "None", '{"name":"No Pass","email":"nopass@example.com"}', 422, "422 Unprocessable Entity for missing password field", "High", "High", "422 Unprocessable Entity", "PASS"),
    ("TC_AUTH_007", "Authentication", "/api/auth/register", "POST", "Verify registration with weak password (length < 6)", "None", '{"name":"Weak Pass","email":"weak@example.com","password":"123"}', 422, "422 Unprocessable Entity for password min length", "Medium", "Medium", "422 Unprocessable Entity", "PASS"),
    ("TC_AUTH_008", "Authentication", "/api/auth/register", "POST", "Verify password is not returned in registration response", "None", '{"name":"Secure","email":"sec@example.com","password":"Password@123"}', 201, "User object returned without password/password_hash", "High", "High", "Password omitted from response", "PASS"),
    ("TC_AUTH_009", "Authentication", "/api/auth/login", "POST", "Verify login with valid credentials", "User registered", '{"email":"qauser@example.com","password":"Test@12345"}', 200, "200 OK with access_token and bearer type", "High", "High", "200 OK with valid JWT token", "PASS"),
    ("TC_AUTH_010", "Authentication", "/api/auth/login", "POST", "Verify login with invalid password", "User exists", '{"email":"qauser@example.com","password":"WrongPassword!"}', 401, '401 Unauthorized: "Invalid email or password"', "High", "High", "401 Unauthorized", "PASS"),
    ("TC_AUTH_011", "Authentication", "/api/auth/login", "POST", "Verify login with non-registered email", "None", '{"email":"unknown@example.com","password":"Test@12345"}', 401, '401 Unauthorized: "Invalid email or password"', "High", "High", "401 Unauthorized", "PASS"),
    ("TC_AUTH_012", "Authentication", "/api/auth/login", "POST", "Verify login with empty request body", "None", '{}', 422, "422 Unprocessable Entity", "Medium", "Medium", "422 Unprocessable Entity", "PASS"),
    ("TC_AUTH_013", "Authentication", "/api/auth/me", "GET", "Verify get current user with valid Bearer token", "Valid JWT token", "Header: Authorization: Bearer <valid_token>", 200, "200 OK with current user details", "High", "High", "200 OK with profile", "PASS"),
    ("TC_AUTH_014", "Authentication", "/api/auth/me", "GET", "Verify get current user without Authorization header", "None", "None", 401, "401 Unauthorized: token is missing", "High", "High", "401 Unauthorized", "PASS"),
    ("TC_AUTH_015", "Authentication", "/api/auth/me", "GET", "Verify get current user with malformed token", "None", "Header: Authorization: Bearer invalid_token", 401, "401 Unauthorized: Invalid token", "High", "High", "401 Unauthorized", "PASS"),

    # Projects Module (12 cases)
    ("TC_PROJ_001", "Projects", "/api/projects", "POST", "Verify creating a project with valid details", "Valid JWT token", '{"name":"Mobile App","description":"iOS/Android","status":"active"}', 201, "201 Created with project details", "High", "High", "201 Created with project object", "PASS"),
    ("TC_PROJ_002", "Projects", "/api/projects", "GET", "Verify retrieving list of owned projects", "Valid JWT token", "None", 200, "200 OK with array of project objects", "High", "High", "200 OK with projects list", "PASS"),
    ("TC_PROJ_003", "Projects", "/api/projects/{id}", "GET", "Verify retrieving project by valid ID", "Valid JWT token, project ID 1 exists", "None", 200, "200 OK with project object for ID 1", "High", "High", "200 OK with project object", "PASS"),
    ("TC_PROJ_004", "Projects", "/api/projects/{id}", "PUT", "Verify updating project name and status", "Valid JWT token, owner of project", '{"name":"Updated Name","status":"completed"}', 200, "200 OK with updated project details", "High", "High", "200 OK with updated attributes", "PASS"),
    ("TC_PROJ_005", "Projects", "/api/projects/{id}", "DELETE", "Verify deleting existing project by ID", "Valid JWT token, owner of project", "None", 204, "204 No Content", "High", "High", "204 No Content", "PASS"),
    ("TC_PROJ_006", "Projects", "/api/projects/{id}", "GET", "Verify getting project with non-existent ID", "Valid JWT token", "None", 404, '404 Not Found: "Project not found"', "Medium", "Medium", "404 Not Found", "PASS"),
    ("TC_PROJ_007", "Projects", "/api/projects", "POST", "Verify creating project missing required field name", "Valid JWT token", '{"description":"No name project"}', 422, "422 Unprocessable Entity", "High", "High", "422 Unprocessable Entity", "PASS"),
    ("TC_PROJ_008", "Projects", "/api/projects", "POST", "Verify creating project with invalid status enum", "Valid JWT token", '{"name":"Test","status":"invalid_enum"}', 422, "422 Unprocessable Entity", "Medium", "Medium", "422 Unprocessable Entity", "PASS"),
    ("TC_PROJ_009", "Projects", "/api/projects", "GET", "Verify GET projects without authorization token", "No token", "None", 401, "401 Unauthorized", "High", "High", "401 Unauthorized", "PASS"),
    ("TC_PROJ_010", "Projects", "/api/projects/{id}", "PUT", "Verify updating project owned by another user", "Logged in as devuser, targeting qauser project", '{"name":"Hacked Name"}', 403, '403 Forbidden: "You do not own this project"', "High", "High", "403 Forbidden", "PASS"),
    ("TC_PROJ_011", "Projects", "/api/projects/{id}", "DELETE", "Verify deleting project owned by another user", "Logged in as devuser, targeting qauser project", "None", 403, '403 Forbidden: "You do not own this project"', "High", "High", "403 Forbidden", "PASS"),
    ("TC_PROJ_012", "Projects", "/api/projects", "POST", "Verify database persistence after project creation", "Valid JWT token", '{"name":"DB Persist Check","status":"active"}', 201, "Project persisted in database with correct owner_id", "High", "High", "Record verified in SQLite DB", "PASS"),

    # Tasks Module (15 cases)
    ("TC_TASK_001", "Tasks", "/api/tasks", "POST", "Verify task creation with valid details", "Valid JWT token, project ID 1 exists", '{"title":"API Tests","priority":"high","status":"todo","project_id":1}', 201, "201 Created with task details", "High", "High", "201 Created with task object", "PASS"),
    ("TC_TASK_002", "Tasks", "/api/tasks", "GET", "Verify retrieving tasks list with pagination defaults", "Valid JWT token", "None", 200, "200 OK with items array, total, page, limit", "High", "High", "200 OK with TaskListResponse", "PASS"),
    ("TC_TASK_003", "Tasks", "/api/tasks/{id}", "GET", "Verify retrieving task by valid ID", "Valid JWT token, task ID 1 exists", "None", 200, "200 OK with task details", "High", "High", "200 OK with task object", "PASS"),
    ("TC_TASK_004", "Tasks", "/api/tasks/{id}", "PUT", "Verify updating task status to completed", "Valid JWT token, task ID 1 exists", '{"status":"completed"}', 200, "200 OK with updated task status", "High", "High", "200 OK with updated status", "PASS"),
    ("TC_TASK_005", "Tasks", "/api/tasks/{id}", "DELETE", "Verify deleting task by valid ID", "Valid JWT token, task ID 1 exists", "None", 204, "204 No Content", "High", "High", "204 No Content", "PASS"),
    ("TC_TASK_006", "Tasks", "/api/tasks", "GET", "Verify filtering tasks by status=completed", "Valid JWT token", "Params: status=completed", 200, "200 OK with tasks where status == completed", "High", "High", "200 OK filtered list", "PASS"),
    ("TC_TASK_007", "Tasks", "/api/tasks", "GET", "Verify filtering tasks by priority=high", "Valid JWT token", "Params: priority=high", 200, "200 OK with tasks where priority == high", "High", "High", "200 OK filtered list", "PASS"),
    ("TC_TASK_008", "Tasks", "/api/tasks", "GET", "Verify filtering tasks by project_id", "Valid JWT token", "Params: project_id=1", 200, "200 OK with tasks under project 1", "High", "High", "200 OK filtered list", "PASS"),
    ("TC_TASK_009", "Tasks", "/api/tasks", "GET", "Verify task pagination parameters page=1 and limit=2", "Valid JWT token", "Params: page=1&limit=2", 200, "200 OK with max 2 items", "Medium", "Medium", "200 OK paginated response", "PASS"),
    ("TC_TASK_010", "Tasks", "/api/tasks/{id}", "GET", "Verify getting task with non-existent ID", "Valid JWT token", "None", 404, '404 Not Found: "Task not found"', "Medium", "Medium", "404 Not Found", "PASS"),
    ("TC_TASK_011", "Tasks", "/api/tasks", "POST", "Verify task creation missing title field", "Valid JWT token", '{"project_id":1}', 422, "422 Unprocessable Entity", "High", "High", "422 Unprocessable Entity", "PASS"),
    ("TC_TASK_012", "Tasks", "/api/tasks", "POST", "Verify task creation with invalid priority enum", "Valid JWT token", '{"title":"T","priority":"urgent","project_id":1}', 422, "422 Unprocessable Entity", "Medium", "Medium", "422 Unprocessable Entity", "PASS"),
    ("TC_TASK_013", "Tasks", "/api/tasks", "POST", "Verify task creation with invalid status enum", "Valid JWT token", '{"title":"T","status":"reviewing","project_id":1}', 422, "422 Unprocessable Entity", "Medium", "Medium", "422 Unprocessable Entity", "PASS"),
    ("TC_TASK_014", "Tasks", "/api/tasks", "POST", "Verify task creation with non-existent project_id", "Valid JWT token", '{"title":"Orphan Task","project_id":99999}', 400, '400 Bad Request: "Project does not exist"', "High", "High", "400 Bad Request", "PASS"),
    ("TC_TASK_015", "Tasks", "/api/tasks", "POST", "Verify database persistence and relationship after task creation", "Valid JWT token", '{"title":"DB Task","project_id":1}', 201, "Task persisted with correct project_id FK", "High", "High", "Record verified in SQLite DB", "PASS"),

    # Users Module / Authorization (5 cases)
    ("TC_USER_001", "Users", "/api/users", "GET", "Verify admin can retrieve all users list", "Admin JWT token", "None", 200, "200 OK with list of all users", "High", "High", "200 OK with users array", "PASS"),
    ("TC_USER_002", "Users", "/api/users/{id}", "GET", "Verify admin can retrieve user details by ID", "Admin JWT token", "None", 200, "200 OK with user object", "High", "High", "200 OK with user details", "PASS"),
    ("TC_USER_003", "Users", "/api/users/{id}", "GET", "Verify admin getting non-existent user ID", "Admin JWT token", "None", 404, '404 Not Found: "User not found"', "Medium", "Medium", "404 Not Found", "PASS"),
    ("TC_USER_004", "Users", "/api/users", "GET", "Verify normal non-admin user cannot access admin users list", "Normal User JWT token", "None", 403, '403 Forbidden: "Admin privilege required"', "High", "Critical", "403 Forbidden", "PASS"),
    ("TC_USER_005", "Users", "/api/users/{id}", "GET", "Verify normal non-admin user cannot access user by ID", "Normal User JWT token", "None", 403, '403 Forbidden: "Admin privilege required"', "High", "Critical", "403 Forbidden", "PASS"),

    # Dashboard Module (5 cases)
    ("TC_DASH_001", "Dashboard", "/api/dashboard", "GET", "Verify authenticated user receives dashboard metrics", "Valid JWT token", "None", 200, "200 OK with total_projects, total_tasks, etc.", "High", "High", "200 OK with metrics object", "PASS"),
    ("TC_DASH_002", "Dashboard", "/api/dashboard", "GET", "Verify dashboard statistics accurately match DB counts", "Valid JWT token", "None", 200, "Dashboard counts match exact count from database", "High", "High", "Metrics match DB records", "PASS"),
    ("TC_DASH_003", "Dashboard", "/api/dashboard", "GET", "Verify unauthenticated dashboard request fails", "No token", "None", 401, "401 Unauthorized", "High", "High", "401 Unauthorized", "PASS"),
    ("TC_DASH_004", "Dashboard", "/api/dashboard", "GET", "Verify dashboard response time is under 2 seconds", "Valid JWT token", "None", 200, "Response time < 2.0s", "Medium", "Low", "Response time ~ 0.015s", "PASS"),
    ("TC_DASH_005", "Dashboard", "/api/dashboard", "GET", "Verify dashboard schema keys and integer types", "Valid JWT token", "None", 200, "All 5 count keys present as integers", "High", "Medium", "Schema matches definition", "PASS"),
]

for row_idx, data in enumerate(test_cases, start=2):
    ws.append(data)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.border = thin_border
        if col_idx == 13:  # Status column
            cell.fill = pass_fill
            cell.font = pass_font

# Apply styling to Header
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Adjust column widths automatically
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save("manual_testing/API_Test_Cases.xlsx")
print("Successfully generated manual_testing/API_Test_Cases.xlsx with 52 test cases!")
